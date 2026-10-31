#!/usr/bin/env python3
"""Parse ODFW controlled hunt draw reports into HuntScout Pro DrawRecord JSON.

Inputs (all produced by draw_download.py in <raw_dir>):
  <year>_<series>_points.xlsx  ODFW "Preference Point Draw Report" (100/200/400/600/700 series)
  <year>_<series>_choice.xlsx  ODFW "Applicants by Hunt Choice" (sheep, goat, premium L/M/N)
  regs_<year>.txt              pdftotext -layout of the Oregon Big Game Hunting Regulations
                               (hunt names, resident/nonresident tag numbers and previous-year
                               resident/nonresident applications for sheep and goat; premium
                               hunt lists)

Usage: python -I draw_parse.py <raw_dir> <out_dir>
Writes <out_dir>/<species>.json. Nothing is estimated: values absent from the agency
documents are left out (null). Odds are computed drawn/applicants and flagged oddsComputed.
"""
import json
import os
import re
import sys

import openpyxl

YEARS = [2022, 2023, 2024, 2025, 2026]
POINT_SERIES = {"buck": "100", "elk": "200", "pronghorn": "400", "antlerless": "600", "springbear": "700"}
PREMIUM = {"premium_buck": "L", "premium_elk": "M", "premium_pronghorn": "N"}
# regs text files available (2025 booklet could not be retrieved)
REGS_YEARS = [2022, 2023, 2024, 2026]
# Hunt names missing from every retrieved booklet, as printed in ODFW's draft 2026 tag table
# (raw: draft_2026_Rocky_Mtn_Goat_DRAFT.pdf, row "916A Central Cascades").
DRAFT_NAMES = {"goat": {"916A": "Central Cascades"}}


def sid(year, series, kind):
    return f"d-or-{year}-{series}-{kind}"


def regs_sid(year):
    return f"d-or-regs-{year}"


# ---------------------------------------------------------------- helpers
def code_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    s = str(v).strip().replace(",", "")
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    return None


def odds(drawn, apps):
    if drawn is None or not apps:
        return None
    return round(100.0 * drawn / apps, 2)


def method_of(name):
    n = name or ""
    if re.search(r"Archery/Muzzleloader ONLY", n, re.I):
        return "other", "Archery/Muzzleloader only"
    if re.search(r"Traditional Bow", n, re.I):
        return "archery", "Traditional bow"
    if re.search(r"\bBow\b", n):
        return "archery", "Bow"
    if re.search(r"Muzzleloader", n, re.I):
        return "muzzleloader", "Muzzleloader"
    return None, None


def is_whitetail(name):
    return bool(re.search(r"whitetail|white-tail", name or "", re.I))


# Deer hunts whose regulations bag limit is white-tailed deer only ("One white-tailed deer",
# "One antlerless white-tailed deer", "One white-tailed deer buck with visible antler"), even
# when the hunt name does not say "Whitetail" (e.g. 152M Grande Ronde Muzzleloader, 657A Flora).
# Mixed bag limits ("One white-tailed or black-tailed buck", "One white-tailed buck ... or ...")
# are not matched. Added by the verification pass.
_WT_CODE_RE = re.compile(r"^\s*((?:1|6)\d\d[A-Z0-9]*|[A-Z]{2}(?:1|6)\d\d[A-Z0-9-]*)\s")
_WT_BAG_RE = re.compile(r"One (?:antlerless )?white-tailed deer(?! or)(?:\s+buck with visible antler)?\s*(?:$|\s{2,})")


def parse_regs_whitetail(path):
    lines = open(path, encoding="utf-8", errors="replace").read().splitlines()
    out = set()
    for i, line in enumerate(lines):
        m = _WT_CODE_RE.match(line)
        if not m:
            continue
        seg = line[m.end():]
        if "white-tailed" in seg:
            after = seg.split("white-tailed", 1)[1][:30].split()
            if _WT_BAG_RE.search(seg) and "or" not in after:
                out.add(m.group(1))
            continue
        if re.search(r"One |deer|buck|elk|antelope|pronghorn", seg):
            continue
        # bag limit printed on the line above/below (wrapped table cell)
        for j in (i - 1, i + 1):
            if 0 <= j < len(lines) and not _WT_CODE_RE.match(lines[j]) and _WT_BAG_RE.search(lines[j]) and len(lines[j].strip()) < 60:
                out.add(m.group(1))
    return out


# ---------------------------------------------------------------- points report
def parse_points(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and code_str(r[0]) == "Hunt Number")
    hdr = [str(c).strip() if c is not None else "" for c in rows[h]]
    assert hdr[2].startswith("Tags") and hdr[8] == "Total Drawn", hdr
    assert "Points" in hdr[12] and hdr[13] == "Applicants", hdr
    hunts, order, cur = {}, [], None
    for r in rows[h + 1:]:
        if r is None or all(c is None for c in r):
            continue
        code = code_str(r[0])
        if code is not None:
            if not re.fullmatch(r"[A-Z0-9][A-Z0-9-]*", code):
                continue  # titles / footers
            if cur is None or cur["code"] != code or num(r[2]) != cur["tags"]:
                if code in hunts:
                    raise ValueError(f"{path}: hunt {code} appears twice")
                cur = {
                    "code": code, "name": str(r[1]).strip() if r[1] is not None else "",
                    "tags": num(r[2]), "ra": num(r[3]), "rd": num(r[4]), "na": num(r[5]),
                    "nd": num(r[6]), "ta": num(r[7]), "td": num(r[8]), "pts": [],
                }
                hunts[code] = cur
                order.append(code)
        if cur is None:
            continue
        p = num(r[12])
        if p is None:
            continue
        cur["pts"].append((p, num(r[13]), num(r[14]), num(r[15]), num(r[16]), num(r[17])))
    return [hunts[c] for c in order]


def min_points(rows, appi, drawi, total_apps, total_drawn):
    """Lowest point class with a draw, read from the point table.

    If every first-choice applicant of the class drew (drawn == applicants), the lowest
    point class present drew. Otherwise the lowest class with preference-phase draws > 0.
    """
    present = [r for r in rows if (r[appi] or 0) > 0]
    if not present or not total_drawn:
        return None
    if total_apps is not None and total_drawn == total_apps:
        return min(r[0] for r in present)
    drew = [r[0] for r in rows if (r[drawi] or 0) > 0]
    return min(drew) if drew else None


# ---------------------------------------------------------------- choice report
def _raw_sheet_rows(path):
    """Minimal reader for an XLSX whose workbook.xml openpyxl rejects (2024 buck choice file)."""
    import zipfile
    from xml.etree import ElementTree as ET
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    z = zipfile.ZipFile(path)
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")))
    rows = []
    for row in ET.fromstring(z.read("xl/worksheets/sheet1.xml")).iter(f"{{{ns['m']}}}row"):
        cells = {}
        for c in row.findall("m:c", ns):
            ref = re.match(r"([A-Z]+)", c.get("r")).group(1)
            col = 0
            for ch in ref:
                col = col * 26 + ord(ch) - 64
            v = c.find("m:v", ns)
            val = None
            if c.get("t") == "s" and v is not None:
                val = shared[int(v.text)]
            elif c.get("t") == "inlineStr":
                val = "".join(t.text or "" for t in c.iter(f"{{{ns['m']}}}t"))
            elif v is not None:
                val = float(v.text) if re.search(r"[.eE]", v.text) else int(v.text)
            cells[col - 1] = val
        rows.append(tuple(cells.get(i) for i in range(max(cells) + 1)) if cells else ())
    return rows


def parse_choice(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if wb.worksheets:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    else:
        rows = _raw_sheet_rows(path)
    h = next(i for i, r in enumerate(rows) if r and code_str(r[0]) == "Hunt Number")
    hdr = [str(c).strip() if c is not None else "" for c in rows[h]]
    assert hdr[1] == "1st Choice", hdr
    out = {}
    for r in rows[h + 1:]:
        code = code_str(r[0]) if r else None
        if not code or not re.fullmatch(r"[A-Z0-9][A-Z0-9-]*", code):
            continue
        out[code] = num(r[1])
    return out


# ---------------------------------------------------------------- regulations
SEASON_RE = re.compile(r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)\.? \d{1,2} - "
                       r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)\.? \d{1,2})")


def parse_regs_sheep_goat(path, year):
    """Return {code: {name, season, resTags, nrTags, prevResApps, prevNrApps}} for 500/900 series."""
    lines = open(path, encoding="utf-8", errors="replace").read().split("\n")
    out = {}
    i = 0
    while i < len(lines):
        L = lines[i]
        if re.search(r"ANY LEGAL WEAPON CONTROLLED (CALIFORNIA BIGHORN SHEEP|ROCKY MTN BIGHORN SHEEP|ROCKY MTN GOAT) HUNTS", L):
            # find header line with "<year> RESIDENT TAGS"
            j = i + 1
            while j < len(lines) and f"{year} RESIDENT TAGS" not in lines[j]:
                j += 1
            hdr = lines[j]
            c_rt = hdr.index(f"{year} RESIDENT TAGS") + len(f"{year} RESIDENT TAGS") / 2
            c_nt = hdr.index(f"{year} NONRESIDENT TAGS") + len(f"{year} NONRESIDENT TAGS") / 2
            # applications columns: "RESIDENT"/"NONRESIDENT" words on the line above or the header
            # itself, followed by APPLICATIONS; use positions of the two APPLICATIONS/RESIDENT tokens
            above = lines[j - 1] + " " * 400
            below = lines[j + 1] + " " * 400
            ap = [m.start() for m in re.finditer(r"APPLICATIONS", hdr)]
            if len(ap) < 2:
                ap = [m.start() for m in re.finditer(r"APPLICATIONS", below)]
            c_ra = ap[0] + 6
            c_na = ap[1] + 6
            k = j + 1
            prev_text = None
            while k < len(lines):
                s = lines[k]
                if re.search(r"ANY LEGAL WEAPON|OREGON BIG GAME HUNTING REGULATIONS\s*\d*$|^\s*\d+\s+20\d\d OREGON", s):
                    break
                m = re.match(r"^(\s*)(5\d\d[A-Z]?\d?|9\d\d[A-Z]?\d?)\s+(.*)$", s)
                if m and SEASON_RE.search(s):
                    code = m.group(2)
                    sm = SEASON_RE.search(s)
                    name = s[m.end(2):sm.start()].strip()
                    if not name:
                        nxt = lines[k + 1].strip() if k + 1 < len(lines) else ""
                        name = " ".join(x for x in (prev_text, nxt) if x)
                        k += 1  # consume continuation line
                    name = re.sub(r"\s+", " ", name).replace("*", "").strip()
                    rec = {"name": name, "season": sm.group(1), "resTags": None, "nrTags": None,
                           "prevResApps": None, "prevNrApps": None, "prevNrNA": False}
                    for t in re.finditer(r"[\d,]+|N/A", s[sm.end():]):
                        pos = sm.end() + t.start() + len(t.group()) / 2
                        cols = {"resTags": c_rt, "nrTags": c_nt, "prevResApps": c_ra, "prevNrApps": c_na}
                        key = min(cols, key=lambda c: abs(cols[c] - pos))
                        if t.group() == "N/A":
                            if key == "prevNrApps":
                                rec["prevNrNA"] = True
                            continue
                        rec[key] = num(t.group())
                    out[code] = rec
                    prev_text = None
                else:
                    st = s.strip()
                    prev_text = st if st and not re.search(r"\d", st) and len(st) < 40 else None
                k += 1
            i = k
            continue
        i += 1
    return out


def parse_regs_premium(path):
    """Return {code: name} for premium L/M/N hunts listed in the regulations."""
    txt = open(path, encoding="utf-8", errors="replace").read()
    out = {}
    for sec in ("DEER", "PRONGHORN ANTELOPE", "ELK"):
        m = re.search(rf"ANY LEGAL WEAPON {sec} HUNTS\n", txt)
        if not m:
            continue
        block = txt[m.end(): m.end() + 6000]
        letter = {"DEER": "L", "PRONGHORN ANTELOPE": "N", "ELK": "M"}[sec]
        for line in block.split("\n"):
            if "ANY LEGAL WEAPON" in line:
                break
            for mm in re.finditer(rf"\b({letter}[A-Z]{{0,2}}\d{{2}}[A-Z]?)\s+(.+?)(?=\s{{3,}}|$)", line):
                out[mm.group(1)] = mm.group(2).replace("*", "").strip()
    return out


MON = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)\.?"
HUNT_LINE_RE = re.compile(
    rf"^\s*([1-7]\d\d[A-Z0-9]*|[A-Z]{{2}}[16]0\d[A-Z0-9-]*)\s+(.+?)\s+({MON} \d{{1,2}} - {MON} \d{{1,2}})\s+([\d,]+)\s+([\d,]+|NEW|N/A)?\s*$",
    re.M)


def parse_regs_hunt_lines(path):
    """Controlled hunt table rows of the regulations booklet that fit on one line:
    {code: {season, tags, prevApps}} where tags = '<year> tags' and prevApps = previous-year
    first-choice applicants as printed. Hunts with split seasons/wrapped rows are skipped."""
    txt = open(path, encoding="utf-8", errors="replace").read()
    out = {}
    for m in HUNT_LINE_RE.finditer(txt):
        code = m.group(1)
        if code in out:
            out[code] = None  # ambiguous: printed twice
            continue
        out[code] = {"season": m.group(3), "tags": num(m.group(4)), "prevApps": num(m.group(5)) if m.group(5) else m.group(5)}
    return {k: v for k, v in out.items() if v}


# ---------------------------------------------------------------- build
def build(raw, outdir):
    recs = {k: [] for k in ("elk", "mule-deer", "whitetail", "pronghorn", "bear", "sheep", "goat")}
    regs = {y: os.path.join(raw, f"regs_{y}.txt") for y in REGS_YEARS}
    hunt_lines = {y: parse_regs_hunt_lines(regs[y]) for y in REGS_YEARS}
    wt_regs = {y: parse_regs_whitetail(regs[y]) for y in REGS_YEARS}
    # 2025 booklet not retrieved: carry a 2024 white-tailed-only hunt forward only when the 2025
    # report lists the same hunt code with the same hunt name (ODFW's 2025 changes were limited to
    # season dates and tag numbers, ODFW news release 2024-09-13).
    names24 = {}
    for ser in ("buck", "antlerless"):
        for hh in parse_points(os.path.join(raw, f"2024_{ser}_points.xlsx")):
            names24[hh["code"]] = hh["name"]

    def wt_basis(year, code, name):
        if is_whitetail(name):
            return "name"
        if year in wt_regs and code in wt_regs[year]:
            return f"regs{year}"
        if year == 2025 and code in wt_regs[2024] and names24.get(code) == name:
            return "regs2024"
        return None

    phase_note = ("Counts are first-choice applicants/draws from ODFW's Preference Point Draw Report. "
                  "byPoints drawn counts (and minPoints) cover only the 75% preference-point phase; tags "
                  "awarded in the 25% random phase are in the hunt totals but are not broken out by point level.")

    for series, prefix in POINT_SERIES.items():
        for y in YEARS:
            path = os.path.join(raw, f"{y}_{series}_points.xlsx")
            for h in parse_points(path):
                if series == "elk":
                    sp = "elk"
                elif series == "pronghorn":
                    sp = "pronghorn"
                elif series == "springbear":
                    sp = "bear"
                else:
                    wtb = wt_basis(y, h["code"], h["name"])
                    sp = "whitetail" if wtb else "mule-deer"
                byp = []
                for p, a, ra, rd, na, nd in sorted(h["pts"], key=lambda t: t[0]):
                    byp.append({"points": p,
                                "resident": {"applicants": ra, "drawn": rd} if ra is not None and rd is not None else None,
                                "nonresident": {"applicants": na, "drawn": nd} if na is not None and nd is not None else None})
                tot_rows = [(p, a, (rd or 0) + (nd or 0)) for p, a, ra, rd, na, nd in h["pts"]]
                res_rows = [(p, ra, rd) for p, a, ra, rd, na, nd in h["pts"]]
                nr_rows = [(p, na, nd) for p, a, ra, rd, na, nd in h["pts"]]
                rec = {
                    "species": sp, "year": y, "huntCode": h["code"], "units": [], "description": h["name"],
                }
                meth, mlab = method_of(h["name"])
                if meth:
                    rec["method"] = meth
                    rec["methodLabel"] = mlab
                if series == "buck":
                    rec["sex"] = "male"
                hl = hunt_lines.get(y, {}).get(h["code"])
                season_from_regs = bool(hl and hl["tags"] == h["tags"])
                if season_from_regs:
                    rec["seasonDates"] = hl["season"]
                rec["choice"] = "first"
                rec["resident"] = {"applicants": h["ra"], "drawn": h["rd"], "oddsPct": odds(h["rd"], h["ra"]),
                                   "oddsComputed": True, "minPoints": min_points(res_rows, 1, 2, h["ra"], h["rd"])}
                rec["nonresident"] = {"applicants": h["na"], "drawn": h["nd"], "oddsPct": odds(h["nd"], h["na"]),
                                      "oddsComputed": True, "minPoints": min_points(nr_rows, 1, 2, h["na"], h["nd"])}
                rec["total"] = {"quota": h["tags"], "applicants": h["ta"], "drawn": h["td"], "oddsPct": odds(h["td"], h["ta"]),
                                "oddsComputed": True, "minPoints": min_points(tot_rows, 1, 2, h["ta"], h["td"])}
                notes = []
                # ---- consistency of the agency's own table (report as published, flag problems)
                pr = sum(t[2] or 0 for t in h["pts"])
                pn = sum(t[4] or 0 for t in h["pts"])
                p1d = sum((t[3] or 0) + (t[5] or 0) for t in h["pts"])
                for key, apps, drawn, lab in (("resident", h["ra"], h["rd"], "resident"), ("nonresident", h["na"], h["nd"], "nonresident")):
                    if apps is not None and drawn is not None and drawn > apps:
                        rec[key]["applicants"] = None
                        rec[key]["oddsPct"] = None
                        rec[key]["minPoints"] = None
                        notes.append(f"ODFW's report prints {apps} {lab} applicants but {drawn} {lab} drawn for this hunt (total applicants {h['ta']}); "
                                     f"the {lab} applicant count is left blank.")
                if byp and h["pts"] and (p1d > (h["td"] or 0) or pr > 1.5 * (h["ra"] or 0) + 5):
                    for k in ("resident", "nonresident", "total"):
                        rec[k]["minPoints"] = None
                    notes.append(f"ODFW's point-class table for this hunt sums to {pr} resident and {pn} nonresident applicants with {p1d} preference-phase draws, "
                                 f"inconsistent with the hunt totals ({h['ra']} resident, {h['na']} nonresident, {h['td']} drawn); point-class rows omitted.")
                    byp = []
                elif byp and h["pts"] and (pr != h["ra"] or pn != h["na"]):
                    notes.append(f"As published, the point-class rows sum to {pr} resident / {pn} nonresident applicants vs {h['ra']} / {h['na']} in the hunt totals.")
                if byp:
                    rec["byPoints"] = byp
                if series == "buck":
                    notes.append("100 series controlled buck deer hunt.")
                elif series == "antlerless":
                    notes.append("600 series controlled antlerless deer hunt.")
                elif series == "springbear":
                    notes.append("700 series controlled spring black bear hunt.")
                elif series == "elk":
                    notes.append("200 series controlled elk hunt.")
                elif series == "pronghorn":
                    notes.append("400 series controlled pronghorn hunt.")
                notes.append(phase_note)
                notes.append("Tags = tags authorized; tags left after first choice go to 2nd-5th choice applicants (not shown).")
                if series == "springbear" and y == 2026:
                    notes.append("ODFW's 2026 spring bear report labels the point-class column 'Points Spent'.")
                if season_from_regs:
                    notes.append(f"Season dates from the {y} Oregon Big Game Hunting Regulations ({regs_sid(y)}).")
                if sp == "mule-deer" and series in ("buck", "antlerless"):
                    notes.append("Filed under mule deer because neither the hunt name nor the regulations bag limit makes it a white-tailed-deer-only hunt; western Oregon deer hunts are for black-tailed deer.")
                if sp == "whitetail" and wtb and wtb.startswith("regs"):
                    ry = int(wtb[4:])
                    if ry == y:
                        notes.append(f"Filed under whitetail because the {y} regulations ({regs_sid(y)}) give this hunt a white-tailed-deer-only bag limit.")
                    else:
                        notes.append(f"Filed under whitetail because the {ry} regulations ({regs_sid(ry)}) give hunt {h['code']} ('{h['name']}') a white-tailed-deer-only bag limit; "
                                     f"the {y} booklet was not retrieved and ODFW's {y} changes were limited to season dates and tag numbers.")
                rec["note"] = " ".join(notes)
                rec["sourceId"] = sid(y, series, "points")
                recs[sp].append(rec)

    # ---- premium hunts (L deer, M elk, N pronghorn): random, one tag per hunt
    prem_names = {y: parse_regs_premium(regs[y]) for y in REGS_YEARS}
    for series, letter in PREMIUM.items():
        sp = {"L": "mule-deer", "M": "elk", "N": "pronghorn"}[letter]
        for y in YEARS:
            ch = parse_choice(os.path.join(raw, f"{y}_{series}_choice.xlsx"))
            for code, first in ch.items():
                notes = ["Premium hunt (random draw, no preference points; residents and nonresidents have equal odds). "
                         "ODFW's Applicants by Hunt Choice report gives first-choice applicants only (no resident/nonresident split, no drawn count)."]
                rec = {"species": sp, "year": y, "huntCode": code, "units": []}
                quota = None
                if y in prem_names and code in prem_names[y]:
                    rec["description"] = prem_names[y][code]
                    quota = 1
                    rec["seasonDates"] = "Aug. 1 - Nov. 30"
                    notes.append(f"Hunt name, quota and season from the {y} Oregon Big Game Hunting Regulations ({regs_sid(y)}): "
                                 "'One tag is available for each hunt listed.' Season: Aug. 1 - Nov. 30.")
                else:
                    alt = next((yy for yy in (2026, 2024, 2023, 2022) if code in prem_names[yy]), None)
                    nm = prem_names[alt][code] if alt else None
                    rec["description"] = nm if nm else f"Premium hunt {code}"
                    if alt:
                        notes.append(f"Hunt name taken from the {alt} regulations ({regs_sid(alt)}); the {y} booklet was not retrieved, so the {y} tag number is not shown.")
                    else:
                        notes.append(f"Hunt name not found in the retrieved regulations; tag number not shown.")
                rec["method"] = "all"
                rec["methodLabel"] = "Any legal weapon"
                rec["sex"] = "either"
                rec["choice"] = "first"
                rec["total"] = {"quota": quota, "applicants": first}
                if letter == "L":
                    notes.append("Premium deer tags are good for any species of deer legal to hunt in that unit/hunt area.")
                rec["note"] = " ".join(notes)
                rec["sourceId"] = sid(y, series, "choice")
                recs[sp].append(rec)

    # ---- bighorn sheep (500) and mountain goat (900): random, no points
    sg_regs = {y: parse_regs_sheep_goat(regs[y], y) for y in REGS_YEARS}
    for series in ("sheep", "goat"):
        for y in YEARS:
            ch = parse_choice(os.path.join(raw, f"{y}_{series}_choice.xlsx"))
            for code, first in ch.items():
                cur = sg_regs.get(y, {}).get(code)
                nxt = sg_regs.get(y + 1, {}).get(code)
                notes = [("Bighorn sheep" if series == "sheep" else "Rocky Mountain goat") +
                         " controlled hunt: random draw, no preference points. Total applicants = first-choice applicants from ODFW's Applicants by Hunt Choice report; ODFW does not publish drawn counts for this series."]
                rec = {"species": series, "year": y, "huntCode": code, "units": []}
                name = cur["name"] if cur else None
                if not name:
                    alt = next((yy for yy in (2026, 2024, 2023, 2022) if code in sg_regs[yy]), None)
                    if alt:
                        name = sg_regs[alt][code]["name"]
                        notes.append(f"Hunt name taken from the {alt} regulations ({regs_sid(alt)}).")
                if not name and code in DRAFT_NAMES.get(series, {}):
                    name = DRAFT_NAMES[series][code]
                    notes.append("Hunt name from ODFW's draft 2026 900-series tag table (d-or-draft-2026-goat); "
                                 "the draft lists this hunt with 0 tags for 2026 (rotation with CTWS).")
                rec["description"] = name or f"Hunt {code}"
                if not name:
                    notes.append("Hunt name not found in the retrieved regulations.")
                rec["method"] = "all"
                rec["methodLabel"] = "Any legal weapon"
                if series == "sheep":
                    rec["sex"] = "female" if "Ewe" in (name or "") else "male"
                else:
                    rec["sex"] = "either"
                if cur:
                    rec["seasonDates"] = cur["season"]
                rec["choice"] = "first"
                res = {}
                nr = {}
                if cur:
                    res["quota"] = cur["resTags"]
                    if cur["nrTags"] is not None:
                        nr["quota"] = cur["nrTags"]
                    else:
                        notes.append(f"The {y} regulations list no nonresident tags for this hunt.")
                    notes.append(f"Resident/nonresident tag numbers and season from the {y} Oregon Big Game Hunting Regulations ({regs_sid(y)}).")
                else:
                    notes.append(f"{y} tag numbers not shown: the {y} regulations booklet was not retrieved.")
                if nxt and y + 1 == 2023:
                    notes.append("Resident/nonresident application counts are not shown for 2022: the '2022 applications' columns of the 2023 regulations "
                                 "count all hunt choices (they sum to the report's all-choice total), not first-choice applicants.")
                elif nxt:
                    if nxt["prevResApps"] is not None:
                        res["applicants"] = nxt["prevResApps"]
                    if nxt["prevNrApps"] is not None:
                        nr["applicants"] = nxt["prevNrApps"]
                    notes.append(f"Resident/nonresident application counts are the '{y} applications' columns of the {y + 1} regulations ({regs_sid(y + 1)})"
                                 + ("; nonresident applications shown as N/A there." if nxt.get("prevNrNA") else "."))
                rec["resident"] = res or None
                rec["nonresident"] = nr or None
                rec["total"] = {"applicants": first}
                rec["note"] = " ".join(notes)
                rec["sourceId"] = sid(y, series, "choice")
                recs[series].append(rec)

    os.makedirs(outdir, exist_ok=True)
    for sp, rows in recs.items():
        rows.sort(key=lambda r: (r["year"], r["huntCode"]))
        p = os.path.join(outdir, f"{sp}.json")
        with open(p, "w") as f:
            json.dump(rows, f, indent=1, ensure_ascii=False)
        print(sp, len(rows), os.path.getsize(p))


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
